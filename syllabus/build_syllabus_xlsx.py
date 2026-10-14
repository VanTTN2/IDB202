"""Build the IDB201 syllabus workbook from the university template.

Usage (from the repository root):
    pip install openpyxl
    python syllabus/build_syllabus_xlsx.py

Input : syllabus/template/Syllabus_Temp.xlsx   (the official template, unchanged)
Output: syllabus/IDB201_Syllabus.xlsx

All syllabus content lives in the data structures below, so to revise the syllabus
you edit this file and run it again.
"""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "template" / "Syllabus_Temp.xlsx"
OUTPUT = ROOT / "IDB201_Syllabus.xlsx"

# ---------------------------------------------------------------------------
# 1. General information (sheet "Syllabus", column C)
# ---------------------------------------------------------------------------
GENERAL = {
    "Document type": "SYLLABUS",
    "Program": "UNDERGRADUATE PROGRAM – Bachelor of Computer Science.",
    "Decision No.": "",
    "Course Name": "Introduction to Databases_Giới thiệu về cơ sở dữ liệu",
    "Course Code": "IDB201",
    "Leaning-Teaching Method": "Offline",
    "No of credits": "3",
    "Degree Level": "Bachelor",
    "Time Allocation": "45h (60 sessions) contact hours + 1h final exam + 104h self-study",
    "Pre-requisite": "None (semester 1).",
    "Description": (
        "This course aims to help students:\n"
        "- Understand the formal foundations of the relational model: relations as sets, and keys and constraints "
        "as logical statements;\n"
        "- Design databases with the Entity-Relationship model and convert them into relational schemas;\n"
        "- Write queries in relational algebra, relational calculus and SQL, and reason about what each language "
        "can express;\n"
        "- Program the database server in T-SQL (functions, stored procedures, triggers and cursors), and compare "
        "row-by-row with set-based processing;\n"
        "- Apply functional-dependency theory and its algorithms to normalize a design, and prove that a "
        "decomposition is correct;\n"
        "- Analyze the cost of storage, indexing and query-processing algorithms;\n"
        "- Explain the system internals that keep data correct: transactions, concurrency control and recovery;\n"
        "- Answer research-style questions and carry out small experiments on database topics."
    ),
    "Student's tasks": (
        "- Students must attend at least 80% of contact slots in order to be accepted to the final examination.\n"
        "- Read the assigned textbook sections BEFORE each theory session.\n"
        "- For each chapter with constructivist questions (all chapters except C5, T-SQL programming), choose ONE of its three constructivist questions and submit a short written answer "
        "on the LMS within one week of the discussion session. The answers are graded as part of the lab and "
        "constructivist-question score (4% in total).\n"
        "- Complete all 5 labs, including the 'Investigate' part, and submit them on time.\n"
        "- Complete the group project (one assignment for the whole course): a database designed through the "
        "course and a small study answering research questions about it, delivered as an academic report, "
        "reproducible SQL scripts and a presentation.\n"
        "- Complete the guided self-study readings listed in the schedule (abstract topics such as relational "
        "calculus, bag semantics of SQL, statistics estimation and recovery algorithms); they are assessed in the "
        "progress tests and the final exam.\n"
        "- Python (accessing SQL Server from Python; Python versions of the course algorithms) is optional "
        "self-study. It is not taught in class and is not required in the labs.\n"
        "- Keep all code, scripts and measurements in a version-controlled repository so that results are "
        "reproducible.\n"
        "- Disclose any use of generative AI tools in submitted work; undisclosed use is treated as a breach of "
        "academic integrity.\n"
        "- Use laptop in class only for learning purpose.\n"
        "- Promptly access to the FU FLM at https://flm.fpt.edu.vn/ for up-to-date course information."
    ),
    "Tools": (
        "- Microsoft SQL Server (Developer/Express)\n"
        "- SQL Server Management Studio (SSMS)\n"
        "- draw.io (diagrams.net) for ER diagrams\n"
        "- RelaX relational algebra calculator (online, optional)\n"
        "- Git / GitHub for reproducible code and experiments (optional)\n"
        "- Python 3 with pyodbc and pandas; Visual Studio Code or PyCharm (optional, self-study only)"
    ),
    "Note": (
        "1) On-going Assessment:                 50%\n"
        "- 5 labs + 9 CQ answers:                10%\n"
        "- Progress test 1:                          5%\n"
        "- Progress test 2:                          5%\n"
        "- 1 Assignment (group project):      30%\n"
        "2) 1 Final Exam:                            50%\n"
        "3) Final Result                             100%\n"
        "Completion Criteria:\n"
        "1) Every on-going assessment component > 0\n"
        "2) Final Exam Score >= 5 & Final Result >= 5"
    ),
    "Min GPA to pass": "5",
    "Scoring scale": "10",
    "Approved date": "",
}

# ---------------------------------------------------------------------------
# 2. Materials
# ---------------------------------------------------------------------------
# (description, purpose, ISBN, type, note, author, publisher, published date, edition)
MATERIALS = [
    ("Database System Concepts", "textbook", "978-0-07-802215-9", "hardcopy",
     "[DSC] Main textbook. The schedule follows its chapters and cites sections and pages; Ch. 27 (online) is used for self-study",
     "Abraham Silberschatz, Henry F. Korth, S. Sudarshan", "McGraw-Hill Education", "2020", "7th"),
    ("Database Management Systems", "reference", "978-0-07-246563-1", "hardcopy",
     "[DMS] Secondary textbook for alternative explanations and extra exercises", "Raghu Ramakrishnan, Johannes Gehrke",
     "McGraw-Hill", "2003", "3rd"),
    ("Transact-SQL reference", "reference", "", "online",
     "https://learn.microsoft.com/sql/t-sql/", "Microsoft", "Microsoft", "", ""),
]

# ---------------------------------------------------------------------------
# 3. CLOs
# ---------------------------------------------------------------------------
CLOS = [
    ("CLO1", "Explain the fundamental concepts of database systems (data models, schemas and instances, the "
             "three-schema architecture, data independence, DBMS components) and relate them to the research "
             "contributions that introduced them."),
    ("CLO2", "Design a database from requirements: build an Entity-Relationship model, justify design alternatives, "
             "and map it to a relational schema whose keys and integrity constraints are specified precisely, "
             "including as logical statements."),
    ("CLO3", "Formulate queries in relational algebra and relational calculus, prove simple equivalences between "
             "expressions, and explain the limits of their expressive power."),
    ("CLO4", "Use SQL to define, populate, modify and query relational databases (joins, aggregation, subqueries, "
             "views, parameterized queries), program the database server in T-SQL (user-defined functions, stored "
             "procedures, triggers, cursors), and explain query results through SQL semantics (three-valued logic; "
             "bag semantics through guided self-study)."),
    ("CLO5", "Analyze and improve relational schemas with functional dependency theory: compute closures, candidate "
             "keys and minimal covers, prove basic properties, decompose to 3NF/BCNF with lossless-join and "
             "dependency-preservation guarantees."),
    ("CLO6", "Evaluate physical design and query execution choices (storage organization, indexes, join algorithms, "
             "execution plans) using I/O cost models and controlled experiments."),
    ("CLO7", "Analyze concurrent executions for conflict serializability and recoverability, and justify "
             "concurrency-control and recovery mechanisms (two-phase locking, isolation levels, write-ahead logging)."),
    ("CLO8", "Conduct a small, reproducible research study on a database topic: review primary literature, formulate "
             "a research question or hypothesis, design and carry out an experiment or implementation, and report "
             "the results in an academic report and presentation."),
]

# CLO -> PLO mapping for the 13 PLOs of BCS_AD (source: docs/sources/BCS_AD_13_PLO.xlsx)
PLO_COUNT = 13
PLO_MAP = {
    "CLO1": [2],
    "CLO2": [2, 5],
    "CLO3": [2, 5],
    "CLO4": [2, 8],
    "CLO5": [2, 5, 8],
    "CLO6": [2, 5, 9],
    "CLO7": [2],
    "CLO8": [5, 7, 10, 11],
}

# ---------------------------------------------------------------------------
# 4. Schedule: (topic, CLO, ITU, student's materials, lecturer's materials, student's task, lecturer's task)
#    ITU: I = Introduce, T = Teach, U = Utilize
#    The course follows the chapter order of the main textbook [DSC]:
#    Silberschatz, Korth & Sudarshan, Database System Concepts, 7th ed., 2020.
#    Page numbers come from the book's table of contents.
#    Course chapters (C1–C10) and their DSC chapters:
#      C1 Introduction ............................ DSC 1
#      C2 Relational model and relational algebra . DSC 2 (+ online Ch. 27, self-study)
#      C3 Introduction to SQL ..................... DSC 3
#      C4 Intermediate and advanced SQL ........... DSC 4, 5.1, 5.4
#      C5 Programming the server in T-SQL ......... DSC 5.2, 5.3 (+ T-SQL reference)
#      C6 Database design with the E-R model ...... DSC 6
#      C7 Relational database design .............. DSC 7
#      C8 Storage and indexing .................... DSC 12, 13, 14
#      C9 Query processing and optimization ....... DSC 15, 16
#      C10 Transactions, concurrency control and recovery ... DSC 17, 18, 19 (recovery mainly self-study)
# ---------------------------------------------------------------------------
def S(topic, clo, itu, smat, lmat, stask, ltask):
    return (topic, clo, itu, smat, lmat, stask, ltask)


def theory(topic, clo, itu, ch, reading, self_study=""):
    task = f"Read {reading} before class"
    if self_study:
        task += f"\nSelf-study: {self_study}"
    return S(topic, clo, itu,
             f"- Slides: Chapter {ch}\n- Textbook: {reading}",
             f"- Syllabus IDB201\n- Slides: Chapter {ch}\n- Textbook [DSC]",
             task,
             f"Teach Chapter {ch}")


def lab(topic, clo, n, reading, extra=""):
    return S(topic, clo, "U",
             f"- Lab {n} handout{extra}\n- Textbook: {reading}",
             f"- Lab {n} handout and sample solution",
             f"Do Lab {n}; submit before the end of the session",
             f"Guide Lab {n}; review and grade submissions")


def review(topic, clo, smat, lmat, stask, ltask):
    return S(topic, clo, "U", smat, lmat, stask, ltask)


SCHEDULE = [
    # ---- C1 Introduction (DSC Ch. 1) -------------------------------------------------------------
    theory("Course introduction: syllabus, assessment, research orientation\n"
           "1.1 Database-system applications\n1.2 Purpose of database systems\n1.3 View of data",
           "CLO1", "I", 1, "DSC Ch. 1, §1.1–1.3, pp. 1–12"),
    theory("1.4 Database languages\n1.5 Database design\n1.6 Database engine\n1.7 Database and application architecture",
           "CLO1", "T", 1, "DSC Ch. 1, §1.4–1.7, pp. 13–23",
           "DSC §1.8–1.9 (users and administrators; history of database systems), pp. 24–28"),
    S("Discussion of the Chapter 1 constructivist questions\nTool set-up (not graded): install SQL Server and SSMS; load the sample university database",
      "CLO1", "T\nU", "- Slides: Chapter 1\n- Set-up guide\n- Textbook: DSC Appendix A, pp. 1287–1298",
      "- Slides: Chapter 1\n- Set-up guide",
      "Discuss the Chapter 1 constructivist questions; choose one and submit an answer within one week\nInstall the tools and load the sample database",
      "Lead the discussion; help students with the set-up"),
    # ---- C2 Relational model and relational algebra (DSC Ch. 2; online Ch. 27) ------------------
    theory("2.1 Structure of relational databases\n2.2 Database schema\n2.3 Keys\n2.4 Schema diagrams\n"
           "The relational model as mathematics: relations as sets, keys and constraints as logic",
           "CLO2, CLO3", "T", 2, "DSC Ch. 2, §2.1–2.4, pp. 37–46"),
    theory("2.5 Relational query languages\n2.6 The relational algebra: select, project, union, set difference, "
           "Cartesian product, rename",
           "CLO3", "T", 2, "DSC Ch. 2, §2.5–2.6, pp. 47–57"),
    theory("2.6 The relational algebra (cont.): joins, intersection, assignment, equivalent queries; "
           "division and aggregation as extended operators",
           "CLO3", "T", 2, "DSC Ch. 2, §2.6, pp. 48–57",
           "Relational calculus (tuple and domain) and Codd's theorem, DSC online Ch. 27 – guided reading with worksheet"),
    S("Assignment (group project) introduction: project brief, teams of 3–4, milestones and rubric; "
      "how to state a research question, design a small experiment and keep it reproducible",
      "CLO8", "I", "- Assignment brief and rubric", "- Assignment brief, rubric and topic list",
      "Form teams; choose an application domain; read the brief",
      "Present the assignment; approve teams and domains"),
    review("Guided exercises: relational algebra queries (checked with RelaX); algebra and calculus proofs; "
           "review of the relational-calculus self-study worksheet",
           "CLO3", "- Slides: Chapter 2\n- Exercise sheet\n- Relational-calculus worksheet (DSC online Ch. 27)\n"
           "- Textbook: DSC Ch. 2 Exercises, pp. 60–62",
           "- Exercise and worksheet solutions", "Submit the self-study worksheet; solve the algebra and proof exercises",
           "Check the worksheet; discuss common mistakes"),
    # ---- C3 Introduction to SQL (DSC Ch. 3) ------------------------------------------------------
    theory("3.1 Overview of SQL\n3.2 SQL data definition\n3.3 Basic structure of SQL queries",
           "CLO4", "T", 3, "DSC Ch. 3, §3.1–3.3, pp. 65–78"),
    theory("3.4 Additional basic operations\n3.5 Set operations\n3.6 Null values and three-valued logic",
           "CLO4", "T", 3, "DSC Ch. 3, §3.4–3.6, pp. 79–90"),
    theory("3.7 Aggregate functions\n3.8 Nested subqueries",
           "CLO4", "T", 3, "DSC Ch. 3, §3.7–3.8, pp. 91–107"),
    theory("3.9 Modification of the database\nTranslating SQL queries into relational algebra",
           "CLO3, CLO4", "T", 3, "DSC Ch. 3, §3.9, pp. 108–113",
           "Bag (multiset) semantics of SQL: duplicates, ALL vs. DISTINCT, and which algebraic laws fail for bags "
           "(DSC Ch. 3, §3.3–3.5 and §3.7, with a worksheet)"),
    lab("Lab 1 (part 1): single-table queries, joins, NULL behaviour", "CLO4", 1, "DSC Ch. 3 Exercises, pp. 115–123"),
    lab("Lab 1 (part 2): aggregation, subqueries, set operations; equivalent formulations of the same query",
        "CLO3, CLO4", 1, "DSC Ch. 3 Exercises, pp. 115–123"),
    # ---- C4 Intermediate and advanced SQL (DSC Ch. 4, §5.1, §5.4) -------------------------------
    theory("4.1 Join expressions\n4.2 Views", "CLO4", "T", 4, "DSC Ch. 4, §4.1–4.2, pp. 125–142"),
    theory("4.3 Transactions (preview)\n4.4 Integrity constraints\n4.5 SQL data types and schemas",
           "CLO2, CLO4", "T", 4, "DSC Ch. 4, §4.3–4.5, pp. 143–163",
           "DSC §4.7 Authorization, pp. 165–172"),
    theory("4.6 Index definition in SQL\n5.4 Recursive queries: transitive closure and the limits of relational algebra",
           "CLO3, CLO4", "T", 4, "DSC §4.6, p. 164; §5.4, pp. 213–218",
           "DSC §5.5 Advanced aggregation features (ranking, windowing), pp. 219–230"),
    theory("5.1 Accessing SQL from a programming language: concepts (API, cursors, prepared statements)\n"
           "SQL injection and parameterized queries in T-SQL",
           "CLO4", "T", 4, "DSC §5.1, pp. 183–197",
           "DSC §9.8 Application security, pp. 437–446; Python access to SQL Server (pyodbc, pandas), optional"),
    lab("Lab 2: DDL, integrity constraints, views and parameterized queries", "CLO2, CLO4", 2,
        "DSC Ch. 4 Exercises, pp. 176–179; §5.1"),
    # ---- C5 Programming the database server in T-SQL (DSC §5.2–5.3; T-SQL reference) ----------
    theory("5.2 Functions and procedures (1): T-SQL batches, variables and control flow (IF, WHILE); "
           "scalar and table-valued user-defined functions",
           "CLO4", "T", 5, "DSC §5.2 (within pp. 198–212); T-SQL reference: control-of-flow language, CREATE FUNCTION"),
    theory("5.2 Functions and procedures (2): stored procedures with input and output parameters; "
           "error handling with TRY...CATCH; transactions inside a procedure (preview of Chapter 10)",
           "CLO4", "T", 5, "DSC §5.2 (within pp. 198–212); T-SQL reference: CREATE PROCEDURE, TRY...CATCH",
           "Dynamic SQL with sp_executesql and its injection risk (T-SQL reference; links back to session 18)"),
    theory("5.3 Triggers: AFTER and INSTEAD OF triggers; the inserted and deleted tables; enforcing business rules "
           "that keys and CHECK constraints cannot express; when not to use triggers",
           "CLO2, CLO4", "T", 5, "DSC §5.3 (within pp. 198–212); T-SQL reference: CREATE TRIGGER"),
    theory("Cursors: from the cursor concept of §5.1 to T-SQL cursors (DECLARE, OPEN, FETCH, CLOSE); "
           "row-by-row versus set-based processing, and why set-based SQL is usually faster",
           "CLO4", "T", 5, "DSC §5.1 (cursors, revisited) and §5.2 (loops over query results); "
                           "T-SQL reference: DECLARE CURSOR, FETCH"),
    lab("Lab 3 (part 1): user-defined functions and stored procedures with parameters and error handling "
        "on the university database", "CLO4", 3, "DSC Ch. 5 Exercises on §5.2; T-SQL reference"),
    lab("Lab 3 (part 2): triggers for business rules and audit logs; a cursor solution versus a set-based "
        "solution of the same task (Investigate)", "CLO2, CLO4, CLO8", 3, "DSC Ch. 5 Exercises on §5.2–5.3"),
    review("Progress test 1 (Chapters 1–5)\nReview exercises", "CLO1, CLO2, CLO3, CLO4",
           "- Slides: Chapters 1–5\n- DSC Ch. 1–4, Ch. 5 §5.1–5.4", "- Progress test 1",
           "Review Chapters 1–5; take Progress test 1", "Run and grade Progress test 1; review the answers"),
    # ---- C6 Database design with the E-R model (DSC Ch. 6) --------------------------------------
    theory("6.1 Overview of the design process\n6.2 The Entity-Relationship model\n6.3 Complex attributes",
           "CLO2", "T", 6, "DSC Ch. 6, §6.1–6.3, pp. 241–251"),
    theory("6.4 Mapping cardinalities\n6.5 Primary key\n6.6 Removing redundant attributes in entity sets",
           "CLO2", "T", 6, "DSC Ch. 6, §6.4–6.6, pp. 252–263"),
    theory("6.7 Reducing E-R diagrams to relational schemas", "CLO2", "T", 6, "DSC Ch. 6, §6.7, pp. 264–270",
           "DSC §6.8 Extended E-R features, pp. 271–278"),
    theory("6.9 Entity-relationship design issues; limits of the E-R model",
           "CLO2", "T", 6, "DSC Ch. 6, §6.9, pp. 279–284",
           "DSC §6.10–6.11 Alternative notations (crow's foot, UML) and other aspects of design, pp. 285–291"),
    lab("Lab 4: E-R design, reduction to relational schemas and implementation in SQL Server", "CLO2", 4,
        "DSC Ch. 6 Exercises, pp. 294–299"),
    # ---- C7 Relational database design (DSC Ch. 7) ----------------------------------------------
    theory("7.1 Features of good relational designs: redundancy and anomalies\n7.2 Decomposition using functional dependencies",
           "CLO5", "T", 7, "DSC Ch. 7, §7.1–7.2, pp. 303–312"),
    theory("7.3 Normal forms: BCNF and 3NF", "CLO5", "T", 7, "DSC Ch. 7, §7.3, pp. 313–319"),
    theory("7.4 Functional-dependency theory: closure of a set of FDs, Armstrong's axioms, attribute closure",
           "CLO5", "T", 7, "DSC Ch. 7, §7.4, pp. 320–329"),
    theory("7.4 Functional-dependency theory (cont.): canonical cover, lossless decomposition, dependency preservation; "
           "proofs of soundness", "CLO5", "T", 7, "DSC Ch. 7, §7.4, pp. 320–329"),
    theory("7.5 Algorithms for decomposition: BCNF decomposition and 3NF synthesis", "CLO5", "T", 7,
           "DSC Ch. 7, §7.5, pp. 330–335"),
    theory("7.6 Decomposition using multivalued dependencies (4NF, overview)\n7.8 Atomic domains and first normal form", "CLO5", "T", 7, "DSC Ch. 7, §7.6, §7.8, pp. 336–340, 342",
           "DSC §7.7 More normal forms, p. 341; §7.9 Database-design process, pp. 343–346"),
    review("Guided exercises: FD theory on paper (attribute closure, candidate keys, canonical cover, "
           "BCNF and 3NF decomposition, lossless-join test)",
           "CLO5", "- Slides: Chapter 7\n- Textbook: DSC Ch. 7 Exercises, pp. 353–359", "- Exercise solutions",
           "Solve the exercises; check each decomposition for lossless join and dependency preservation",
           "Guide the exercises; discuss solutions"),
    S("Assignment milestone 1: design review (E-R model, normalized schema with constraints, research questions)",
      "CLO2, CLO5, CLO8", "U",
      "- Assignment brief and rubric", "- Assignment rubric",
      "Submit the design document; present it (5 minutes per team)", "Assess the designs; give feedback"),
    review("Progress test 2 (Chapters 6–7)\nReview exercises", "CLO2, CLO5",
           "- Slides: Chapters 6–7\n- DSC Ch. 6–7", "- Progress test 2",
           "Review Chapters 6–7; take Progress test 2", "Run and grade Progress test 2; review the answers"),
    # ---- C8 Storage and indexing (DSC Ch. 12–14) ------------------------------------------------
    theory("13.1 Database storage architecture\n13.2 File organization\n13.3 Organization of records in files\n"
           "13.5 Database buffer",
           "CLO6", "T", 8, "DSC Ch. 13, §13.1–13.3, pp. 587–601; §13.5, pp. 604–610",
           "DSC §12.1 Physical storage media, pp. 559–561; §12.6 Disk-block access, pp. 577–579; "
           "§13.4 Data-dictionary storage, pp. 602–603"),
    theory("14.1 Indexing: basic concepts\n14.2 Ordered indices\n14.3 B+-tree index files: structure, search, height analysis",
           "CLO6", "T", 8, "DSC Ch. 14, §14.1–14.2, pp. 623–633; §14.3 (structure and queries), pp. 634–649"),
    theory("14.3 B+-tree index files (cont.): insertion; deletion (main idea)\n14.5 Hash indices\n"
           "14.7 Creation of indices", "CLO6", "T", 8,
           "DSC Ch. 14, §14.3 (updates), pp. 634–649; §14.5–14.7, pp. 658–664",
           "DSC §14.4 B+-tree extensions, pp. 650–657; §14.6 Multiple-key access; optional: §13.6 Column-oriented "
           "storage, pp. 611–614; §14.8 Write-optimized index structures (LSM), pp. 665–669"),
    review("Guided exercises: B+-tree insertion by hand; hashing; I/O cost of index lookups versus table scans",
           "CLO6", "- Slides: Chapter 8\n- Textbook: DSC Ch. 14 Exercises, pp. 679–682", "- Exercise solutions",
           "Solve the exercises; compare hand-computed costs with the formulas", "Guide the exercises; discuss solutions"),
    # ---- C9 Query processing and optimization (DSC Ch. 15–16) -----------------------------------
    theory("15.1 Overview of query processing\n15.2 Measures of query cost\n15.3 Selection operation\n"
           "15.4 Sorting (external merge sort, main idea)", "CLO6", "T", 9,
           "DSC Ch. 15, §15.1–15.4, pp. 689–703"),
    theory("15.5 Join operation: nested-loop, block nested-loop, indexed nested-loop, merge join and hash join; "
           "cost comparison", "CLO6", "T", 9,
           "DSC Ch. 15, §15.5, pp. 704–718",
           "DSC §15.6 Other operations, pp. 719–723; §15.7 Evaluation of expressions (pipelining), pp. 724–730"),
    theory("16.1 Overview of query optimization\n16.2 Transformation of relational expressions: equivalence rules\n"
           "Reading execution plans in SQL Server", "CLO6", "T", 9,
           "DSC Ch. 16, §16.1–16.2, pp. 743–756",
           "DSC §16.3 Estimating statistics of expression results, pp. 757–765; optional: §16.4 Choice of "
           "evaluation plans (cost-based join ordering by dynamic programming), pp. 766–777"),
    lab("Lab 5 (part 1): indexes and execution plans on a large table; logical reads with and without indexes; "
        "comparing join algorithms in plans (Investigate)",
        "CLO6, CLO8", 5, "DSC Ch. 14–16 Exercises, pp. 679–682, 736–739, 789–793"),
    # ---- C10 Transactions, concurrency control and recovery (DSC Ch. 17–19) --------------------
    theory("17.1 Transaction concept\n17.2 A simple transaction model\n17.3 Storage structure\n"
           "17.4 Transaction atomicity and durability\n17.10 Transactions as SQL statements",
           "CLO7", "T", 10, "DSC Ch. 17, §17.1–17.4, pp. 799–806; §17.10"),
    theory("17.5 Transaction isolation\n17.6 Serializability: conflict serializability and precedence graphs", "CLO7", "T", 10,
           "DSC Ch. 17, §17.5–17.6, pp. 807–818"),
    theory("17.7 Recoverable and cascadeless schedules\n17.8 Transaction isolation levels\n"
           "17.9 Implementation of isolation levels\n18.1 Lock-based protocols: two-phase locking (main idea, lock point)",
           "CLO7", "T", 10, "DSC Ch. 17, §17.7–17.9, pp. 819–827; Ch. 18, §18.1, pp. 835–848",
           "DSC §18.3 Multiple granularity, pp. 853–856"),
    theory("18.2 Deadlock handling\n18.4 Insert operations, delete operations and predicate reads (phantoms)\n"
           "18.8 Snapshot isolation (overview)", "CLO7", "T", 10,
           "DSC Ch. 18, §18.2, pp. 849–852; §18.4, pp. 857–860; §18.8, pp. 872–879",
           "DSC §18.5 Timestamp-based protocols, pp. 861–865; §18.7 Multiversion schemes, pp. 869–871; "
           "§18.9 Weak levels of consistency in practice, pp. 880–882"),
    lab("Lab 5 (part 2): concurrency experiments with two sessions (isolation levels, dirty reads, phantoms, deadlocks); "
        "serializability of the observed schedules by precedence graph",
        "CLO7, CLO8", 5, "DSC Ch. 17–18 Exercises, pp. 831–833, 899–903"),
    theory("Recovery system (overview): 19.1 Failure classification, 19.3 Recovery and atomicity (write-ahead logging), "
           "19.4 Recovery algorithm (main ideas)\nBriefing for the self-study of Chapter 19",
           "CLO7", "T", 10, "DSC Ch. 19, §19.1, §19.3, pp. 907, 912–921",
           "DSC §19.2 Storage, pp. 908–911; §19.4 Recovery algorithm, pp. 922–925; §19.5–19.6, pp. 926–930; "
           "optional §19.9 ARIES, pp. 941–946"),
    # ---- Assignment (group project), review ----------------------------------------------------------
    S("Assignment milestone 2: progress meeting (experiment method, preliminary results)", "CLO6, CLO7, CLO8", "U",
      "- Assignment brief and rubric", "- Assignment rubric", "Present progress; revise the method",
      "Give feedback on each team's method"),
    S("Assignment: final presentations (part 1)", "CLO2, CLO4, CLO5, CLO6, CLO7, CLO8", "U",
      "- Assignment brief", "- Assignment rubric", "Present the project (10 minutes + 5 minutes Q&A)",
      "Assess presentations and reports"),
    S("Assignment: final presentations (part 2)", "CLO2, CLO4, CLO5, CLO6, CLO7, CLO8", "U",
      "- Assignment brief", "- Assignment rubric", "Present the project; submit the final report and SQL scripts",
      "Assess presentations and reports"),
    review("Course review (part 1): Chapters 1–7, including the self-study readings", "CLO1, CLO2, CLO3, CLO4, CLO5",
           "- Slides: Chapters 1–7\n- Self-study reading list", "- Review questions",
           "Review Chapters 1–7 and the self-study readings", "Answer questions; review key results"),
    review("Course review (part 2): Chapters 8–10, including the self-study of recovery", "CLO6, CLO7",
           "- Slides: Chapters 8–10\n- Self-study reading list", "- Review questions",
           "Review Chapters 8–10 and the self-study readings", "Answer questions; review key results"),
    review("Final-exam practice: analysis-level questions (proofs, algorithm traces, cost estimates, schedules)",
           "CLO1, CLO2, CLO3, CLO4, CLO5, CLO6, CLO7",
           "- Practice question set", "- Practice questions and solutions",
           "Answer the practice questions under exam conditions", "Discuss the answers; clarify the exam rules"),
]
assert len(SCHEDULE) == 60, len(SCHEDULE)

# Self-study plan (guided reading, checked by labs, tests and the final exam)
SELF_STUDY = [
    ("C1", "DSC §1.8–1.9, pp. 24–28", "Users and administrators; history of database systems"),
    ("C2", "DSC online Ch. 27 (tuple and domain relational calculus)", "Abstract; studied with a worksheet reviewed in session 8"),
    ("C3", "DSC Ch. 3, §3.3–3.5, §3.7 (duplicates, ALL vs. DISTINCT, aggregates)", "Bag (multiset) semantics of SQL; studied with a worksheet"),
    ("C4", "DSC §4.7, pp. 165–172; §5.5, pp. 219–230; §9.8, pp. 437–446", "Authorization; ranking and windowing; application security"),
    ("C4", "Python with pyodbc and pandas (optional)", "Accessing SQL Server from Python; not taught in class or required in labs"),
    ("C5", "T-SQL reference: dynamic SQL with sp_executesql", "Injection risk of dynamic SQL inside procedures"),
    ("C6", "DSC §6.8, pp. 271–278; §6.10–6.11, pp. 285–291", "Extended E-R features; alternative notations"),
    ("C7", "DSC §7.7, p. 341; §7.9, pp. 343–346", "More normal forms; the design process"),
    ("C7", "Python versions of the FD algorithms (optional)", "Closure, keys, canonical cover, BCNF decomposition"),
    ("C8", "DSC §12.1, §12.6, §13.4, §14.4, §14.6", "Storage media and disk-block access; data dictionary; B+-tree extensions; multiple-key access"),
    ("C9", "DSC §15.6–15.7, pp. 719–730; §16.3, pp. 757–765", "Other operations and pipelining; statistics estimation (abstract, hard to observe directly)"),
    ("C9", "DSC §16.4, pp. 766–777 (optional)", "Cost-based join ordering by dynamic programming; needs algorithm background from CSD203"),
    ("C10", "DSC §18.3, §18.5, §18.7, §18.9", "Multiple granularity; timestamp and multiversion protocols; weak consistency"),
    ("C10", "DSC Ch. 19, §19.2, §19.4–19.6, pp. 908–911, 922–930", "Recovery algorithms: abstract and hard to reproduce in a lab"),
]

# ---------------------------------------------------------------------------
# 5. Constructivist questions: (session, name, question)
# ---------------------------------------------------------------------------
CQ = [
    (3, "CQ1", "Codd separated the logical view of data from its physical storage. What would be lost if applications had to know how and where each record is stored?"),
    (3, "CQ2", "Why do new data models (XML, document, graph, vector) keep appearing, and why do many of them end up adding SQL-like features?"),
    (3, "CQ3", "When is a database system NOT the right tool? Give a concrete example and justify it."),
    (8, "CQ1", "What would break if relations were allowed to contain duplicate tuples? Consider keys, projection and the meaning of a fact."),
    (8, "CQ2", "Division expresses 'for all' queries. Why is there no basic 'for all' operator, and how does relational calculus express the same idea?"),
    (8, "CQ3", "Relational algebra cannot compute the transitive closure of a relation. Why does a fixed-size expression limit what a language can express?"),
    (12, "CQ1", "Under three-valued logic, 'p OR NOT p' is not always true. Which everyday query mistakes follow from this?"),
    (12, "CQ2", "Why does 'x NOT IN (subquery)' return no rows when the subquery contains NULL, while NOT EXISTS behaves differently?"),
    (12, "CQ3", "COUNT(*) and COUNT(attribute) can return different values on the same table. Why, and what does this show about how aggregates treat NULL?"),
    (18, "CQ1", "Why should integrity rules be declared in the database rather than checked in application code? When is the opposite true?"),
    (18, "CQ2", "A view costs nothing extra at run time. How is that possible, and when would a materialized view be better?"),
    (18, "CQ3", "Why do parameterized queries prevent SQL injection, while escaping user input is considered fragile?"),
    (30, "CQ1", "Why is a grade an attribute of the enrollment relationship and not of STUDENT or SECTION? What goes wrong if it is placed on either entity?"),
    (30, "CQ2", "Is a ternary relationship always equivalent to three binary relationships? Construct a counterexample."),
    (30, "CQ3", "Give a business rule of the university database that cannot be drawn in an E-R diagram. How should a designer record and enforce it?"),
    (37, "CQ1", "Every normal form can be seen as removing one kind of redundancy. What single idea unifies them?"),
    (37, "CQ2", "BCNF decomposition may lose a dependency. When would you accept 3NF instead, and what do you gain?"),
    (37, "CQ3", "Testing whether an attribute is prime is NP-complete. What does this mean for automatic schema design tools?"),
    (44, "CQ1", "Why is a B+-tree node the size of a disk block instead of holding a single key as in a binary search tree?"),
    (44, "CQ2", "When can a full table scan be cheaper than using an index? Design an experiment that shows it."),
    (44, "CQ3", "Could a machine-learning model replace an index? What would you need to measure to decide?"),
    (47, "CQ1", "The optimizer relies on estimates. What happens when the estimates are wrong, and where do the errors come from?"),
    (47, "CQ2", "The same query can often be written with a join, a subquery or a set operation. Should the choice affect performance? Why or why not?"),
    (47, "CQ3", "Why does the number of possible join orders grow so quickly, and why can an optimizer not simply try them all?"),
    (54, "CQ1", "Why is serializability the accepted correctness criterion for concurrent transactions? What weaker criteria do real systems use, and why?"),
    (54, "CQ2", "Why does two-phase locking guarantee serializability? Explain the idea of the lock point."),
    (54, "CQ3", "Why must a log record reach stable storage before the data page it describes (write-ahead logging)?"),
]

# Mark the discussion session of each chapter's constructivist questions in the schedule.
for _session in sorted({q[0] for q in CQ}):
    _t = list(SCHEDULE[_session - 1])
    if "constructivist" not in _t[5]:
        _t[5] += "\nDiscuss this chapter's constructivist questions; choose one and submit an answer within one week"
        _t[6] += "; lead the discussion of the constructivist questions"
    SCHEDULE[_session - 1] = tuple(_t)

# ---------------------------------------------------------------------------
# 6. Grading structure
# (component, type, weight, part, min, duration, CLO, question type, number, scope, how, note, reference)
# ---------------------------------------------------------------------------
PT_NOTE_ = "Instruction and schedules for Progress tests must be presented in the Course Implementation Plan approved by director of the campus.\n\nProgress test must be taken right after the last lectures of required material.\n\nInstructor has responsibility to review the test for students after graded."
GRADING = [
    ("Labs and constructivist questions", "on-going", 10, 14, 0.0001, "Labs: in lab sessions; CQ answers: 1 week after each discussion",
     "CLO1, CLO2, CLO3, CLO4, CLO5, CLO6, CLO7, CLO8",
     "Lab exercises in SQL Server, 'Investigate' mini-reports; short written answers to constructivist questions",
     "5 labs + 9 CQ answers",
     "Labs 1–5 (Chapters 3–10: SQL queries; DDL, constraints and views; T-SQL functions, procedures, triggers and cursors; E-R design; indexing, query plans and concurrency); one constructivist question chosen by the student from each of the 9 chapters that have them (all except C5, T-SQL programming)",
     "labs: in class, by instructor; CQ answers: on the LMS, by instructor with a short rubric",
     "Labs 6% in total (1.2% each). Constructivist-question answers 4% in total (shared equally among the 9 answers; graded on reasoning, use of the textbook and clarity). The 'Investigate' part (hypothesis, experiment, evidence) counts for at least 40% of each lab mark.",
     "Lab"),
    ("Progress test 1", "on-going", 5, 1, 0.0001, "30'", "CLO1, CLO2, CLO3, CLO4",
     "Multiple choices (marked by computer) + short written reasoning", "20 MCQ + 2 written",
     "- cover Chapters 1–5 (DSC Ch. 1–4, §5.1–5.4) and their self-study readings (relational calculus, bag semantics)\n- written items: one relational algebra/calculus proof or query, one SQL semantics or T-SQL reasoning question (NULL and three-valued logic, or the effect of a trigger)",
     "in class, by instructor",
     PT_NOTE_,
     "Progress test"),
    ("Progress test 2", "on-going", 5, 1, 0.0001, "30'", "CLO2, CLO5",
     "Multiple choices (marked by computer) + short written reasoning", "20 MCQ + 2 written",
     "- cover Chapters 6–7 (DSC Ch. 6–7) and their self-study sections\n- written items: one E-R design and reduction task, one FD proof or normalization task",
     "in class, by instructor",
     PT_NOTE_,
     "Progress test"),
    ("Assignment (group project)", "on-going", 30, 1, 0.0001, "Take-home, whole semester; 2 milestones + 15' final presentation",
     "CLO2, CLO3, CLO4, CLO5, CLO6, CLO7, CLO8",
     "Group project (3–4 students): academic report, reproducible SQL scripts and data, presentation", 1,
     "- cover Chapters 2–10\n- each team designs one database for its own domain (E-R model, normalized schema with constraints, SQL queries with correctness arguments, T-SQL procedures and triggers for its business rules) and then answers research questions about it through experiments (triggers versus constraints, cursors versus set-based SQL, indexes and query plans, isolation levels) or through constructivist questions linked to the design",
     "take-home; introduced in session 7; milestone 1 (design review) in session 39; milestone 2 (progress meeting) in session 55; final presentations in sessions 56–57; graded by instructor with rubric",
     "Milestone 1 design review 25%, final report 40% (question, method, results, discussion), reproducibility of scripts and data 10%, presentation and Q&A 20%, peer evaluation 5%. Individual marks may differ within a team based on contribution logs. Plagiarism or undisclosed AI-generated content leads to a score of 0. Instructor has responsibility to review the assignment for students after graded.",
     "Assignment"),
    ("Final exam", "final exam", 50, 1, 5, "60'", "CLO1, CLO2, CLO3, CLO4, CLO5, CLO6, CLO7",
     "Multiple choices\nMarked by Computer", 50,
     "concepts, proofs, algorithms and analysis; all chapters, including the self-study readings (about 20% of items); at least 30% of items at the Analyze level or above",
     "by exam board, using computer",
     "The exam questions must be updated or different at least 70% to the previous ones.",
     "Final exam"),
]
assert sum(g[2] for g in GRADING) == 100
assert next(g[2] for g in GRADING if g[1] == "final exam") == 50


# ---------------------------------------------------------------------------
# Writing helpers
# ---------------------------------------------------------------------------
def copy_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def put(ws, row, col, value, style_from=None, wrap=True):
    cell = ws.cell(row=row, column=col, value=value)
    if style_from is not None:
        copy_style(ws.cell(row=style_from, column=col), cell)
    if wrap:
        al = copy(cell.alignment)
        cell.alignment = Alignment(horizontal=al.horizontal, vertical=al.vertical or "top", wrap_text=True)
    return cell


def fill_syllabus(ws):
    for row in range(2, ws.max_row + 1):
        title = ws.cell(row=row, column=2).value
        if not title:
            continue
        key = title.split("\n")[0].strip()
        if key in GENERAL:
            put(ws, row, 3, GENERAL[key])
    # Row heights for the long text rows.
    for row, height in ((3, 30), (11, 30), (12, 150), (13, 150), (14, 110), (15, 210)):
        ws.row_dimensions[row].height = height


def fill_materials(ws):
    # The template keeps authoring notes below the example rows (from row 9 down).
    # Move them below the material list so that no note is overwritten.
    notes = []
    for row in range(9, ws.max_row + 1):
        for cell in ws[row]:
            if cell.value is not None:
                notes.append((row, cell.column, cell.value, copy(cell._style)))
                cell.value = None
    for i, m in enumerate(MATERIALS):
        row = 2 + i
        put(ws, row, 1, i + 1, style_from=2)
        for col, value in enumerate(m, start=2):
            put(ws, row, col, value, style_from=2)
    offset = max(0, (2 + len(MATERIALS) + 2) - 9)
    for row, col, value, style in notes:
        cell = ws.cell(row=row + offset, column=col, value=value)
        cell._style = style


def fill_clo(ws):
    from openpyxl.styles import Font
    for i, (name, desc) in enumerate(CLOS):
        row = 2 + i
        put(ws, row, 1, i + 1, style_from=2)
        put(ws, row, 2, name, style_from=2)
        put(ws, row, 3, desc, style_from=2)
        ws.row_dimensions[row].height = 45
    ws.cell(row=11, column=1).value = "Mapping of CLOs to PLOs of Curriculum BCS_AD (K22A, 13 PLOs)"
    # Keep only the program's PLO columns in the header (template provides 17).
    for p in range(1, 18):
        ws.cell(row=13, column=1 + p).value = f"PLO{p}" if p <= PLO_COUNT else None
    tick_font = copy(ws["K14"].font)          # template tick: "ü" in Wingdings = check mark
    for i, (name, _) in enumerate(CLOS):
        row = 14 + i
        put(ws, row, 1, name, style_from=14)
        for col in range(2, 22):
            c = ws.cell(row=row, column=col)
            c.value = None
            copy_style(ws.cell(row=14, column=col), c)
        for p in PLO_MAP[name]:
            c = ws.cell(row=row, column=1 + p, value="ü")
            c.font = copy(tick_font)
            c.alignment = Alignment(horizontal="center", vertical="center")


def fill_cq(ws):
    for i, (session, name, question) in enumerate(CQ):
        row = 2 + i
        put(ws, row, 1, i + 1, style_from=2)
        put(ws, row, 2, session, style_from=2)
        put(ws, row, 3, name, style_from=2)
        put(ws, row, 4, question, style_from=2)
    for row in range(2 + len(CQ), ws.max_row + 1):
        for col in range(1, 5):
            ws.cell(row=row, column=col).value = None


def fill_schedule(ws):
    for i, (topic, clo, itu, smat, lmat, stask, ltask) in enumerate(SCHEDULE):
        row = 2 + i
        values = [i + 1, topic, "Offline", clo, itu, smat, lmat, stask, ltask, None, None]
        for col, value in enumerate(values, start=1):
            put(ws, row, col, value, style_from=2)
        ws.row_dimensions[row].height = max(45, 15 * (topic.count("\n") + 2))


def fill_grading(ws):
    # Clear the template's example rows (they can extend below the shorter component list).
    for row in range(2, ws.max_row + 1):
        for col in range(1, 15):
            ws.cell(row=row, column=col).value = None
    for i, g in enumerate(GRADING):
        row = 2 + i
        (comp, typ, weight, part, minv, dur, clo, qtype, num, scope, how, note, ref) = g
        values = [i + 1, comp, typ, weight, part, minv, dur, clo, qtype, num, scope, how, note, ref]
        for col, value in enumerate(values, start=1):
            put(ws, row, col, value, style_from=2)
        ws.row_dimensions[row].height = 150
    total = 2 + len(GRADING)
    put(ws, total, 2, "Total", style_from=2)
    put(ws, total, 4, f"=SUM(D2:D{total - 1})", style_from=2)
    ws.cell(row=total, column=2).font = copy(ws.cell(row=1, column=2).font)


def main():
    wb = openpyxl.load_workbook(TEMPLATE)
    fill_syllabus(wb["Syllabus"])
    fill_materials(wb["Materials"])
    fill_clo(wb["CLO"])
    fill_cq(wb["Constructivist Question"])
    fill_schedule(wb["Schedule"])
    fill_grading(wb["Grading structure"])
    wb.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    main()


# ---------------------------------------------------------------------------
# Markdown export (keeps syllabus/syllabus.md in sync with the workbook)
# ---------------------------------------------------------------------------
def build_markdown(path=ROOT / "syllabus.md"):
    esc = lambda s: str(s).replace("\n", "<br>").replace("|", "\\|")
    L = ["# IDB201 – Introduction to Databases: Course Syllabus", "",
         "> Generated by `syllabus/build_syllabus_xlsx.py` from the same data as "
         "[`IDB201_Syllabus.xlsx`](IDB201_Syllabus.xlsx). **Status: proposal for review by the syllabus council.**", "",
         "## 1. General information", "", "| Item | Details |", "|---|---|"]
    for k in ("Program", "Course Name", "Course Code", "Leaning-Teaching Method", "No of credits", "Degree Level",
              "Time Allocation", "Pre-requisite", "Min GPA to pass", "Scoring scale"):
        L.append(f"| {k.replace('Leaning', 'Learning')} | {esc(GENERAL[k])} |")
    L += ["", "## 2. Description", "", GENERAL["Description"], "",
          "## 3. Student's tasks", "", GENERAL["Student's tasks"], "",
          "## 4. Tools", "", GENERAL["Tools"], "",
          "## 5. Course learning outcomes", "", "| CLO | Description |", "|---|---|"]
    L += [f"| {n} | {esc(d)} |" for n, d in CLOS]
    L += ["", "### CLO–PLO mapping (BCS_AD, 13 PLOs)", "",
          "| CLO | " + " | ".join(f"PLO{p}" for p in range(1, PLO_COUNT + 1)) + " |",
          "|---|" + "---|" * PLO_COUNT]
    for n, _ in CLOS:
        L.append(f"| {n} | " + " | ".join("✓" if p in PLO_MAP[n] else "" for p in range(1, PLO_COUNT + 1)) + " |")
    L += ["", "## 6. Learning materials", "",
          "| # | Material | Purpose | Type | Author | Publisher | Year | Edition | Note |", "|---|---|---|---|---|---|---|---|---|"]
    for i, (desc, purpose, isbn, typ, note, author, pub, year, ed) in enumerate(MATERIALS, 1):
        L.append(f"| {i} | {esc(desc)}{' (ISBN ' + isbn + ')' if isbn else ''} | {purpose} | {typ} | {esc(author)} | {pub} | {year} | {ed} | {esc(note)} |")
    L += ["", "## 7. Assessment", "",
          "| # | Component | Weight | Duration | CLOs | Format | Scope |", "|---|---|---|---|---|---|---|"]
    for i, g in enumerate(GRADING, 1):
        L.append(f"| {i} | {g[0]} | {g[2]}% | {esc(g[5])} | {g[6]} | {esc(g[7])} | {esc(g[9])} |")
    L += ["", "**Completion criteria:** every on-going component > 0; final exam ≥ 5; final result ≥ 5.", "",
          "## 8. Schedule (60 sessions × 45 minutes)", "", "| Session | Topic | CLO | ITU | Student's task |", "|---|---|---|---|---|"]
    for i, s in enumerate(SCHEDULE, 1):
        L.append(f"| {i} | {esc(s[0])} | {esc(s[1])} | {esc(s[2])} | {esc(s[5])} |")
    L += ["", "ITU: I = Introduce, T = Teach, U = Utilize.", "",
          "## 9. Constructivist questions", "", "| Session | Name | Question |", "|---|---|---|"]
    L += [f"| {s} | {n} | {esc(q)} |" for s, n, q in CQ]
    path.write_text("\n".join(L) + "\n", encoding="utf-8")
    print(f"Saved {path}")


if __name__ == "__main__":
    build_markdown()
