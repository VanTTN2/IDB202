"""Build the CLO4 rubric for IDB201 as an Excel workbook and a Markdown file.

Usage:  pip install openpyxl && python docs/rubrics/build_clo4_rubric.py
Output: docs/rubrics/IDB201_CLO4_Rubric.xlsx and docs/rubrics/CLO4_rubric.md
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent

CLO4 = ("Write SQL and T-SQL to define, query and modify relational databases, program server-side logic "
        "(functions, stored procedures, triggers, cursors), and explain query results through SQL semantics "
        "(three-valued logic, bag semantics).")
BLOOM = "Apply"

LEVELS = [  # (name, score range on a 10-point scale)
    ("Excellent", "8.5–10"),
    ("Good", "7.0–8.4"),
    ("Satisfactory", "5.0–6.9"),
    ("Unsatisfactory", "< 5.0"),
]

# (criterion, weight %, evidence, [descriptor per level, from Excellent to Unsatisfactory])
CRITERIA = [
    ("A. Data definition and modification\n(DDL, constraints, DML, views)", 20,
     "Lab 2; group project",
     ["Scripts run without errors. Every key, foreign key, CHECK and default constraint required by the "
      "specification is declared and named. Views return the required data.",
      "Scripts run. Most required constraints are declared; small omissions (e.g. one missing CHECK) do not "
      "break data integrity.",
      "Scripts run after minor fixes. Primary and foreign keys are present, but several business rules are not "
      "enforced.",
      "Scripts fail, or keys and foreign keys are missing, so the data can violate the specification."]),
    ("B. Query formulation\n(joins, aggregation, subqueries, set operations, parameterized queries)", 35,
     "Lab 1, Lab 2; Progress test 1; group project; final exam",
     ["Every query returns exactly the required result, including edge cases (NULLs, empty groups, duplicates). "
      "Constructs are well chosen, and user input is always passed as parameters.",
      "Queries are correct on typical data; at most one edge case is handled incorrectly.",
      "Single-table queries, inner joins and simple aggregation are correct; outer joins, subqueries or set "
      "operations are partly wrong.",
      "Most queries are wrong or do not run."]),
    ("C. Server-side programming in T-SQL\n(user-defined functions, stored procedures, triggers, cursors)", 30,
     "Lab 3; group project",
     ["Functions, procedures and triggers work for all specified cases, handle errors (TRY…CATCH, THROW), and "
      "triggers handle multi-row statements. A cursor is used only when justified, and the set-based "
      "alternative is discussed.",
      "Code works for normal cases; error handling or multi-row handling is incomplete.",
      "A simple function or procedure works; triggers work only for single-row statements; little error "
      "handling.",
      "Code does not compile, or it does not implement the required rule."]),
    ("D. SQL semantics\n(three-valued logic, bag semantics)", 15,
     "Progress test 1 (written item); final exam",
     ["Correctly predicts and justifies results involving NULLs (three-valued logic, NOT IN versus NOT EXISTS) "
      "and duplicates (bags versus sets, UNION versus UNION ALL), with counterexamples.",
      "Predictions are correct, but the justification is partial.",
      "Recognizes that NULLs and duplicates change the result, but the explanation is incomplete or partly "
      "wrong.",
      "Cannot predict or explain the effect of NULLs or duplicates."]),
]
assert sum(c[1] for c in CRITERIA) == 100

SCORING = [
    "CLO4 score of a student = weighted average of the four criteria (weights above), each scored on a "
    "10-point scale from the CLO4-tagged items of Labs 1–3, Progress test 1, the group project and the final exam.",
    "A student achieves CLO4 when the CLO4 score is at least 5.0 (Satisfactory or above).",
    "Course-level target (proposed): at least 70% of students achieve CLO4, and at least 30% reach Good or "
    "Excellent.",
]

FILL = ["C6EFCE", "DDEBF7", "FFEB9C", "FFC7CE"]


def build_xlsx(path):
    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    wb = Workbook()
    ws = wb.active
    ws.title = "CLO4 rubric"
    ws["A1"] = "IDB201 – Rubric for CLO4"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws["A2"] = f"CLO4 ({BLOOM}): {CLO4}"
    ws["A2"].alignment = wrap
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 45

    hdr = ["Criterion", "Weight", "Evidence"] + [f"{n}\n({r})" for n, r in LEVELS]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    for i, (crit, w, ev, desc) in enumerate(CRITERIA):
        r = 5 + i
        vals = [crit, f"{w}%", ev] + desc
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = wrap
            cell.border = border
            if c >= 4:
                cell.fill = PatternFill("solid", fgColor=FILL[c - 4])
        ws.cell(row=r, column=1).font = Font(bold=True)
        ws.row_dimensions[r].height = 120
    r = 5 + len(CRITERIA) + 1
    ws.cell(row=r, column=1, value="Scoring and achievement").font = Font(bold=True)
    for j, t in enumerate(SCORING):
        cell = ws.cell(row=r + 1 + j, column=1, value=f"{j + 1}. {t}")
        ws.merge_cells(start_row=r + 1 + j, start_column=1, end_row=r + 1 + j, end_column=7)
        cell.alignment = wrap
        ws.row_dimensions[r + 1 + j].height = 32
    for col, wdt in zip("ABCDEFG", (30, 9, 22, 34, 34, 34, 34)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A5"
    wb.save(path)


def build_md(path):
    L = ["# IDB201 – Rubric for CLO4", "",
         f"**CLO4 ({BLOOM}):** {CLO4}", "",
         "| Criterion | Weight | Evidence | " + " | ".join(f"{n} ({r})" for n, r in LEVELS) + " |",
         "|---|---|---|" + "---|" * len(LEVELS)]
    for crit, w, ev, desc in CRITERIA:
        L.append(f"| {crit.replace(chr(10), ' ')} | {w}% | {ev} | " + " | ".join(desc) + " |")
    L += ["", "## Scoring and achievement", ""] + [f"{i + 1}. {t}" for i, t in enumerate(SCORING)]
    path.write_text("\n".join(L) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build_xlsx(HERE / "IDB201_CLO4_Rubric.xlsx")
    build_md(HERE / "CLO4_rubric.md")
    print("Saved rubric files in", HERE)
